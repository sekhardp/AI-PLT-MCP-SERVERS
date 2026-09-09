from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Any, Literal

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pydantic import BaseModel, Field


class KPICard(BaseModel):
    label: str = Field(..., description="Metric label")
    value: str = Field(..., description="Metric value (e.g. '$4.2M')")
    change: str | None = Field(None, description="Delta (e.g. '+14.2% YoY')")
    trend: Literal["up", "down", "neutral"] | None = Field("up", description="Trend direction")


class ChartSeries(BaseModel):
    name: str
    values: list[float | int]


class ChartConfig(BaseModel):
    chart_type: Literal["bar", "horizontal_bar", "line", "pie", "doughnut"] = "bar"
    title: str | None = None
    categories: list[str]
    series: list[ChartSeries]


class TableConfig(BaseModel):
    headers: list[str]
    rows: list[list[str]]


class Slide(BaseModel):
    slide_number: int
    layout: Literal[
        "title_slide",
        "kpi_grid",
        "chart_and_bullets",
        "two_column_comparison",
        "table_slide",
        "bullet_cards",
    ]
    title: str
    subtitle: str | None = None
    bullet_points: list[str] = Field(default_factory=list)
    kpi_cards: list[KPICard] = Field(default_factory=list)
    chart: ChartConfig | None = None
    table: TableConfig | None = None
    left_column_title: str | None = None
    left_column_bullets: list[str] = Field(default_factory=list)
    right_column_title: str | None = None
    right_column_bullets: list[str] = Field(default_factory=list)
    sources: list[str] = Field(default_factory=list)


class SlideDeck(BaseModel):
    deck_title: str
    deck_subtitle: str | None = None
    theme: Literal["dark", "light", "midnight", "navy", "emerald"] = "dark"
    author: str | None = "AI Platform MCP Server"
    slides: list[Slide]
    sources_summary: list[str] = Field(default_factory=list)


@dataclass(frozen=True)
class ThemeColors:
    bg: RGBColor
    card_bg: RGBColor
    card_border: RGBColor
    text_primary: RGBColor
    text_secondary: RGBColor
    accent_primary: RGBColor
    accent_secondary: RGBColor
    accent_positive: RGBColor
    accent_negative: RGBColor
    font_title: str = "Helvetica Neue"
    font_body: str = "Calibri"


THEMES: dict[str, ThemeColors] = {
    "dark": ThemeColors(
        bg=RGBColor(15, 23, 42),
        card_bg=RGBColor(30, 41, 59),
        card_border=RGBColor(51, 65, 85),
        text_primary=RGBColor(255, 255, 255),
        text_secondary=RGBColor(148, 163, 184),
        accent_primary=RGBColor(56, 189, 248),
        accent_secondary=RGBColor(129, 140, 248),
        accent_positive=RGBColor(52, 211, 153),
        accent_negative=RGBColor(248, 113, 113),
    ),
    "light": ThemeColors(
        bg=RGBColor(248, 250, 252),
        card_bg=RGBColor(255, 255, 255),
        card_border=RGBColor(226, 232, 240),
        text_primary=RGBColor(15, 23, 42),
        text_secondary=RGBColor(100, 116, 139),
        accent_primary=RGBColor(2, 132, 199),
        accent_secondary=RGBColor(79, 70, 229),
        accent_positive=RGBColor(5, 150, 105),
        accent_negative=RGBColor(220, 38, 38),
    ),
    "midnight": ThemeColors(
        bg=RGBColor(11, 15, 25),
        card_bg=RGBColor(23, 28, 43),
        card_border=RGBColor(44, 53, 76),
        text_primary=RGBColor(241, 245, 249),
        text_secondary=RGBColor(148, 163, 184),
        accent_primary=RGBColor(168, 85, 247),
        accent_secondary=RGBColor(236, 72, 153),
        accent_positive=RGBColor(52, 211, 153),
        accent_negative=RGBColor(244, 63, 94),
    ),
    "navy": ThemeColors(
        bg=RGBColor(10, 25, 47),
        card_bg=RGBColor(17, 34, 64),
        card_border=RGBColor(35, 53, 84),
        text_primary=RGBColor(204, 214, 246),
        text_secondary=RGBColor(136, 146, 176),
        accent_primary=RGBColor(100, 255, 218),
        accent_secondary=RGBColor(87, 203, 255),
        accent_positive=RGBColor(100, 255, 218),
        accent_negative=RGBColor(255, 107, 107),
    ),
    "emerald": ThemeColors(
        bg=RGBColor(6, 44, 34),
        card_bg=RGBColor(12, 60, 47),
        card_border=RGBColor(22, 85, 68),
        text_primary=RGBColor(236, 253, 245),
        text_secondary=RGBColor(167, 243, 208),
        accent_primary=RGBColor(52, 211, 153),
        accent_secondary=RGBColor(251, 191, 36),
        accent_positive=RGBColor(52, 211, 153),
        accent_negative=RGBColor(248, 113, 113),
    ),
}


class PPTXBuilder:
    def __init__(self, theme_name: str = "dark") -> None:
        self.theme = THEMES.get(theme_name.lower(), THEMES["dark"])
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank_layout = self.prs.slide_layouts[6]

    def build_deck(self, deck: SlideDeck) -> io.BytesIO:
        theme_override = deck.theme.lower() if deck.theme else "dark"
        if theme_override in THEMES:
            self.theme = THEMES[theme_override]

        first_slide = deck.slides[0] if deck.slides and deck.slides[0].layout == "title_slide" else None
        if not first_slide:
            self._render_default_title_slide(deck.deck_title, deck.deck_subtitle, deck.author)

        for slide in deck.slides:
            self._render_slide(slide)

        output = io.BytesIO()
        self.prs.save(output)
        output.seek(0)
        return output

    def _apply_background(self, slide_obj: Any) -> None:
        fill = slide_obj.background.fill
        fill.solid()
        fill.fore_color.rgb = self.theme.bg

    def _render_title_slide_content(
        self, slide_obj: Any, title: str, subtitle: str | None, author: str | None
    ) -> None:
        self._apply_background(slide_obj)
        top_bar = slide_obj.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(1.2), Inches(0.08)
        )
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = self.theme.accent_primary
        top_bar.line.color.rgb = self.theme.accent_primary

        tx_box = slide_obj.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(2.2))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = self.theme.font_title
        p.font.size = Pt(40)
        p.font.bold = True
        p.font.color.rgb = self.theme.text_primary

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.name = self.theme.font_body
            p2.font.size = Pt(20)
            p2.font.color.rgb = self.theme.text_secondary
            p2.space_before = Pt(16)

        if author:
            auth_box = slide_obj.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(11.7), Inches(0.8))
            atf = auth_box.text_frame
            ap = atf.paragraphs[0]
            ap.text = f"Prepared by: {author}"
            ap.font.name = self.theme.font_body
            ap.font.size = Pt(13)
            ap.font.color.rgb = self.theme.text_secondary

    def _render_default_title_slide(self, title: str, subtitle: str | None, author: str | None) -> None:
        slide_obj = self.prs.slides.add_slide(self.blank_layout)
        self._render_title_slide_content(slide_obj, title, subtitle, author)

    def _render_header(self, slide_obj: Any, title: str, subtitle: str | None) -> None:
        tx_box = slide_obj.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = self.theme.font_title
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = self.theme.text_primary

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.name = self.theme.font_body
            p2.font.size = Pt(13)
            p2.font.color.rgb = self.theme.accent_primary
            p2.space_before = Pt(4)

    def _render_footer(self, slide_obj: Any, slide_number: int, sources: list[str]) -> None:
        if sources:
            source_box = slide_obj.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(10.0), Inches(0.4))
            stf = source_box.text_frame
            stf.word_wrap = True
            sp = stf.paragraphs[0]
            sp.text = "Sources: " + " | ".join(sources)
            sp.font.name = self.theme.font_body
            sp.font.size = Pt(9)
            sp.font.color.rgb = self.theme.text_secondary

        num_box = slide_obj.shapes.add_textbox(Inches(11.5), Inches(6.8), Inches(1.0), Inches(0.4))
        ntf = num_box.text_frame
        np = ntf.paragraphs[0]
        np.text = f"{slide_number}"
        np.alignment = PP_ALIGN.RIGHT
        np.font.name = self.theme.font_body
        np.font.size = Pt(10)
        np.font.color.rgb = self.theme.text_secondary

    def _render_slide(self, slide: Slide) -> None:
        slide_obj = self.prs.slides.add_slide(self.blank_layout)
        self._apply_background(slide_obj)

        if slide.layout == "title_slide":
            self._render_title_slide_content(slide_obj, slide.title, slide.subtitle, None)
            return

        self._render_header(slide_obj, slide.title, slide.subtitle)
        self._render_footer(slide_obj, slide.slide_number, slide.sources)

        if slide.layout == "kpi_grid":
            self._render_kpi_grid(slide_obj, slide.kpi_cards, slide.bullet_points)
        elif slide.layout == "chart_and_bullets":
            self._render_chart_and_bullets(slide_obj, slide.chart, slide.bullet_points)
        elif slide.layout == "two_column_comparison":
            self._render_two_column_comparison(
                slide_obj,
                slide.left_column_title,
                slide.left_column_bullets,
                slide.right_column_title,
                slide.right_column_bullets,
            )
        elif slide.layout == "table_slide":
            self._render_table_slide(slide_obj, slide.table, slide.bullet_points)
        elif slide.layout == "bullet_cards":
            self._render_bullet_cards(slide_obj, slide.bullet_points)
        else:
            self._render_bullet_cards(slide_obj, slide.bullet_points)

    def _render_kpi_grid(self, slide_obj: Any, kpi_cards: list[KPICard], bullets: list[str]) -> None:
        num_cards = min(len(kpi_cards), 4)
        if num_cards == 0:
            self._render_bullet_cards(slide_obj, bullets)
            return

        total_width = 11.733
        gap = 0.3
        card_w = (total_width - (gap * (num_cards - 1))) / num_cards
        card_h = 1.9
        start_y = 1.8

        for i, card in enumerate(kpi_cards[:num_cards]):
            x = 0.8 + i * (card_w + gap)
            shape = slide_obj.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(start_y), Inches(card_w), Inches(card_h)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = self.theme.card_bg
            shape.line.color.rgb = self.theme.card_border
            shape.line.width = Pt(1)

            tf = shape.text_frame
            tf.word_wrap = True
            p_label = tf.paragraphs[0]
            p_label.text = card.label.upper()
            p_label.font.name = self.theme.font_body
            p_label.font.size = Pt(10)
            p_label.font.bold = True
            p_label.font.color.rgb = self.theme.text_secondary

            p_val = tf.add_paragraph()
            p_val.text = card.value
            p_val.font.name = self.theme.font_title
            p_val.font.size = Pt(26)
            p_val.font.bold = True
            p_val.font.color.rgb = self.theme.text_primary
            p_val.space_before = Pt(4)

            if card.change:
                p_chg = tf.add_paragraph()
                p_chg.text = f"• {card.change}"
                p_chg.font.name = self.theme.font_body
                p_chg.font.size = Pt(11)
                p_chg.font.bold = True
                p_chg.space_before = Pt(4)
                p_chg.font.color.rgb = (
                    self.theme.accent_negative if card.trend == "down" else self.theme.accent_positive
                )

        if bullets:
            b_box = slide_obj.shapes.add_textbox(Inches(0.8), Inches(4.0), Inches(11.733), Inches(2.6))
            btf = b_box.text_frame
            btf.word_wrap = True
            for j, bullet in enumerate(bullets):
                bp = btf.add_paragraph() if j > 0 else btf.paragraphs[0]
                bp.text = f"•  {bullet}"
                bp.font.name = self.theme.font_body
                bp.font.size = Pt(14)
                bp.font.color.rgb = self.theme.text_primary
                bp.space_before = Pt(10)

    def _render_chart_and_bullets(
        self, slide_obj: Any, chart_config: ChartConfig | None, bullets: list[str]
    ) -> None:
        if not chart_config or not chart_config.series:
            self._render_bullet_cards(slide_obj, bullets)
            return

        chart_data = CategoryChartData()
        chart_data.categories = chart_config.categories
        for s in chart_config.series:
            chart_data.add_series(s.name, s.values)

        chart_type_map = {
            "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "horizontal_bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE,
            "pie": XL_CHART_TYPE.PIE,
            "doughnut": XL_CHART_TYPE.DOUGHNUT,
        }
        xl_type = chart_type_map.get(chart_config.chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)
        chart_shape = slide_obj.shapes.add_chart(
            xl_type, Inches(0.8), Inches(1.8), Inches(6.2), Inches(4.7), chart_data
        )
        chart = chart_shape.chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False

        card_shape = slide_obj.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.4), Inches(1.8), Inches(5.1), Inches(4.7)
        )
        card_shape.fill.solid()
        card_shape.fill.fore_color.rgb = self.theme.card_bg
        card_shape.line.color.rgb = self.theme.card_border

        ctf = card_shape.text_frame
        ctf.word_wrap = True
        cp = ctf.paragraphs[0]
        cp.text = "Key Takeaways & Insights"
        cp.font.name = self.theme.font_title
        cp.font.size = Pt(16)
        cp.font.bold = True
        cp.font.color.rgb = self.theme.accent_primary

        for bullet in bullets:
            bp = ctf.add_paragraph()
            bp.text = f"• {bullet}"
            bp.font.name = self.theme.font_body
            bp.font.size = Pt(13)
            bp.font.color.rgb = self.theme.text_primary
            bp.space_before = Pt(12)

    def _render_two_column_comparison(
        self,
        slide_obj: Any,
        left_title: str | None,
        left_bullets: list[str],
        right_title: str | None,
        right_bullets: list[str],
    ) -> None:
        col_w = Inches(5.6)
        col_h = Inches(4.7)
        y = Inches(1.8)

        left_card = slide_obj.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, col_w, col_h)
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = self.theme.card_bg
        left_card.line.color.rgb = self.theme.card_border
        ltf = left_card.text_frame
        ltf.word_wrap = True
        lp = ltf.paragraphs[0]
        lp.text = left_title or "Quantitative Analysis (BigQuery)"
        lp.font.name = self.theme.font_title
        lp.font.size = Pt(16)
        lp.font.bold = True
        lp.font.color.rgb = self.theme.accent_primary

        for b in left_bullets:
            bp = ltf.add_paragraph()
            bp.text = f"• {b}"
            bp.font.name = self.theme.font_body
            bp.font.size = Pt(13)
            bp.font.color.rgb = self.theme.text_primary
            bp.space_before = Pt(12)

        right_card = slide_obj.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), y, col_w, col_h)
        right_card.fill.solid()
        right_card.fill.fore_color.rgb = self.theme.card_bg
        right_card.line.color.rgb = self.theme.card_border
        rtf = right_card.text_frame
        rtf.word_wrap = True
        rp = rtf.paragraphs[0]
        rp.text = right_title or "Qualitative Context & Market Insights (RAG)"
        rp.font.name = self.theme.font_title
        rp.font.size = Pt(16)
        rp.font.bold = True
        rp.font.color.rgb = self.theme.accent_secondary

        for b in right_bullets:
            bp = rtf.add_paragraph()
            bp.text = f"• {b}"
            bp.font.name = self.theme.font_body
            bp.font.size = Pt(13)
            bp.font.color.rgb = self.theme.text_primary
            bp.space_before = Pt(12)

    def _render_table_slide(
        self, slide_obj: Any, table_config: TableConfig | None, bullets: list[str]
    ) -> None:
        if not table_config or not table_config.headers:
            self._render_bullet_cards(slide_obj, bullets)
            return

        num_rows = len(table_config.rows) + 1
        num_cols = len(table_config.headers)
        table_shape = slide_obj.shapes.add_table(
            num_rows, num_cols, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.5)
        )
        table = table_shape.table

        for col_idx, header in enumerate(table_config.headers):
            cell = table.cell(0, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = self.theme.card_border
            p = cell.text_frame.paragraphs[0]
            p.text = header
            p.font.name = self.theme.font_title
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = self.theme.accent_primary

        for row_idx, row_data in enumerate(table_config.rows):
            for col_idx, val in enumerate(row_data):
                if col_idx < num_cols:
                    cell = table.cell(row_idx + 1, col_idx)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = (
                        self.theme.card_bg if row_idx % 2 == 0 else self.theme.bg
                    )
                    p = cell.text_frame.paragraphs[0]
                    p.text = str(val)
                    p.font.name = self.theme.font_body
                    p.font.size = Pt(11)
                    p.font.color.rgb = self.theme.text_primary

    def _render_bullet_cards(self, slide_obj: Any, bullets: list[str]) -> None:
        if not bullets:
            return
        num_cards = min(len(bullets), 4)
        total_width = 11.733
        gap = 0.25
        card_w = (total_width - (gap * (num_cards - 1))) / num_cards

        for i, bullet in enumerate(bullets[:num_cards]):
            x = 0.8 + i * (card_w + gap)
            shape = slide_obj.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8), Inches(card_w), Inches(4.7)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = self.theme.card_bg
            shape.line.color.rgb = self.theme.card_border

            tf = shape.text_frame
            tf.word_wrap = True
            p_num = tf.paragraphs[0]
            p_num.text = f"PILLAR 0{i + 1}"
            p_num.font.name = self.theme.font_title
            p_num.font.size = Pt(11)
            p_num.font.bold = True
            p_num.font.color.rgb = self.theme.accent_primary

            p_body = tf.add_paragraph()
            p_body.text = bullet
            p_body.font.name = self.theme.font_body
            p_body.font.size = Pt(13)
            p_body.font.color.rgb = self.theme.text_primary
            p_body.space_before = Pt(16)
